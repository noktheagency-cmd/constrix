from pathlib import Path
import json, zipfile, xml.etree.ElementTree as ET
from datetime import date, datetime
import openpyxl

root = Path(__file__).resolve().parents[2]
out = Path(__file__).parent
doc = root / 'CONSTRIX_Texniki_Tapsiriq_v1.docx'
with zipfile.ZipFile(doc) as z:
    xml = ET.fromstring(z.read('word/document.xml'))
ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
paragraphs = [''.join(p.itertext()) for p in []]
paragraphs = [''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in xml.findall('.//w:p',ns)]
(out/'requirements.txt').write_text('\n'.join(p for p in paragraphs if p.strip()),encoding='utf-8')
def value(v):
    return v.isoformat() if isinstance(v,(date,datetime)) else v
books=[]
for name in ['alış aktı.xlsx','Nakopitel F2.xlsx','AZCDG-ARTM-STR-BOQ-0001-R01 ISCG.xlsx']:
    wb=openpyxl.load_workbook(root/name,data_only=False)
    cached=openpyxl.load_workbook(root/name,data_only=True)
    sheets=[]
    for ws in wb:
        cells=[]; lines=[]
        for row in ws:
            populated=[]
            for c in row:
                if c.value is not None:
                    cv=cached[ws.title][c.coordinate].value
                    cells.append({'cell':c.coordinate,'value':value(cv),'formula':c.value if c.data_type=='f' else None,'source':value(c.value),'format':c.number_format})
                    populated.append(f'{c.column_letter}={value(c.value)}'+(f' [cached={cv}]' if c.data_type=='f' else ''))
            if populated: lines.append(f'{row[0].row}: '+' | '.join(populated))
        sheets.append({'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,'merged':[str(r) for r in ws.merged_cells.ranges],'cells':cells})
        (out/(name+'.'+ws.title.replace('/','_')+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
    books.append({'file':name,'sheets':sheets})
(out/'workbooks.json').write_text(json.dumps(books,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps([{'file':b['file'],'sheets':[{'name':s['name'],'rows':s['rows'],'columns':s['columns'],'cells':len(s['cells'])} for s in b['sheets']]} for b in books],ensure_ascii=True))
